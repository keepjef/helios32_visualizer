import sys
import os
import argparse
import datetime
from pathlib import Path
import numpy as np
import pandas as pd

# Импорты для рисования нативных графиков в Excel
from openpyxl.chart import LineChart, Reference

PROJECT_ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(PROJECT_ROOT))

from parser.frame_index import FrameIndex
from parser.helios_pcap_parser import HeliosParser

def generate_report(pcap_path, output_path, msop_port, difop_port, roi_min, roi_max):
    print(f"=== LiDAR CLI Report Generator ===")
    print(f"Input file: {pcap_path}")
    print(f"Output file: {output_path}")
    print(f"Ports: MSOP={msop_port}, DIFOP={difop_port}")
    print(f"ROI Min: {roi_min}")
    print(f"ROI Max: {roi_max}\n")

    if not os.path.exists(pcap_path):
        print(f"[ERROR] File not found: {pcap_path}")
        return

    # 1. Индексация файла (подхватит кэш, если он есть)
    index = FrameIndex(pcap_path, msop_port=msop_port, difop_port=difop_port)
    index.load_or_build()

    total_frames = index.frame_count()
    if total_frames == 0:
        print("[ERROR] No 3D frames found in the PCAP file.")
        return

    parser = HeliosParser(pcap_path, index, msop_port=msop_port)
    difop = index.difop_data

    # 2. Анализ статических данных
    rpm = 600
    ret_mode = "Single Return"
    first = {}

    if difop:
        first = difop[0]
        rpm = first.get("mot_spd", 600)
        ret_mode = first.get("return_mode", "Single Return")

    expected_pkts = 0
    if rpm > 0:
        expected_pkts = int(3000 / (rpm / 60)) if ret_mode == "Dual Return" else int(1500 / (rpm / 60))

    frames_data = []
    total_loss = 0
    total_detected_points = 0

    print(f"Total Frames to process: {total_frames}")
    print("Processing frames (extracting 3D points & applying ROI)...")

    # 3. Сканирование всех кадров
    for f_id in range(total_frames):
        percent = int((f_id + 1) / total_frames * 100)
        print(f"\rProgress: [{f_id + 1}/{total_frames}] {percent}%", end="", flush=True)

        stat = index.frame_stats[f_id] if f_id < len(index.frame_stats) else {"msop_count": 0}
        msop = stat["msop_count"]
        loss = expected_pkts - msop
        total_loss += loss

        frame_data = parser.load_frame(f_id)
        det_count = 0

        if frame_data is not None and len(frame_data) > 0:
            xyz = frame_data[:, :3]
            in_box = (
                (xyz[:, 0] >= roi_min[0]) & (xyz[:, 0] <= roi_max[0]) &
                (xyz[:, 1] >= roi_min[1]) & (xyz[:, 1] <= roi_max[1]) &
                (xyz[:, 2] >= roi_min[2]) & (xyz[:, 2] <= roi_max[2])
            )
            det_count = np.sum(in_box)
            total_detected_points += det_count

        frames_data.append({
            "Frame ID": f_id + 1,
            "Expected MSOP Packets": expected_pkts,
            "Received MSOP Packets": msop,
            "Packet Loss": loss,
            "Detection (Points in Box)": det_count
        })

    print("\n\nFrames processed successfully. Generating Excel sheets and Charts...")

    # 4. Формирование датафреймов
    df_frames = pd.DataFrame(frames_data)

    telemetry_data = []
    for item in difop:
        dt = datetime.datetime.fromtimestamp(item["ts"])
        sys_time = dt.strftime('%H:%M:%S.%f')[:-3]
        telemetry_data.append({
            "System Time": sys_time,
            "Hardware Time": item.get("hw_time", ""),
            "Sync Mode": item.get("sync_mode", ""),
            "Sync State": item.get("sync_state", ""),
            "GPS Locks (PPS/RMC/UTC)": f"{item.get('pps_lock', 0)}/{item.get('gprmc_lock', 0)}/{item.get('utc_lock', 0)}",
            "Voltage (V)": item.get("voltage"),           # None вместо "N/A" для графиков
            "Current (A)": item.get("current"),
            "Bot FPGA Temp (°C)": item.get("bot_fpga_temp"),
            "Main Bot Temp (°C)": item.get("main_bot_temp"),
            "Main FPGA Temp (°C)": item.get("main_fpga_temp"),
            "Real RPM": item.get("realtime_rpm")
        })
    df_telem = pd.DataFrame(telemetry_data)

    summary_data = {
        "Metric": [
            "PCAP File Path", "Total Frames", "Total MSOP Packet Loss", "Total Points Detected in ROI",
            "Return Mode", "Target Motor RPM", "LiDAR IP", "MAC Address", "Serial Number",
            "Main Firmware", "Motor Firmware"
        ],
        "Value": [
            pcap_path, total_frames, total_loss, total_detected_points,
            ret_mode, rpm, first.get("lidar_ip", "N/A"), first.get("mac", "N/A"),
            first.get("sn", "N/A"), first.get("top_frm", "N/A"), first.get("mot_frm", "N/A")
        ]
    }
    df_summary = pd.DataFrame(summary_data)

    # 5. Сохранение в Excel и отрисовка ГРАФИКОВ
    try:
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            df_summary.to_excel(writer, sheet_name='Summary', index=False)
            df_frames.to_excel(writer, sheet_name='Frames & Detection', index=False)
            df_telem.to_excel(writer, sheet_name='Telemetry', index=False)
            
            # --- Авто-настройка ширины столбцов ---
            for sheet_name in writer.sheets:
                worksheet = writer.sheets[sheet_name]
                for col in worksheet.columns:
                    max_length = max((len(str(cell.value)) for cell in col if cell.value), default=0)
                    worksheet.column_dimensions[col[0].column_letter].width = (max_length + 2)

            # === ГРАФИКИ ДЛЯ КАДРОВ И ДЕТЕКЦИИ ===
            ws_frames = writer.sheets['Frames & Detection']
            max_row_frames = len(df_frames) + 1
            cat_frames = Reference(ws_frames, min_col=1, min_row=2, max_row=max_row_frames) # Ось X (Frame ID)

            # График 1: Пакеты
            chart_pkts = LineChart()
            chart_pkts.title = "MSOP Packets (Expected vs Received)"
            chart_pkts.y_axis.title = "Packet Count"
            chart_pkts.x_axis.title = "Frame ID"
            chart_pkts.width = 18; chart_pkts.height = 8
            
            data_pkts = Reference(ws_frames, min_col=2, max_col=3, min_row=1, max_row=max_row_frames)
            chart_pkts.add_data(data_pkts, titles_from_data=True)
            chart_pkts.set_categories(cat_frames)
            ws_frames.add_chart(chart_pkts, "G2") # Рисуем в ячейке G2

            # График 2: Детекция
            chart_det = LineChart()
            chart_det.title = "Detection (Points in ROI)"
            chart_det.y_axis.title = "Points Count"
            chart_det.x_axis.title = "Frame ID"
            chart_det.width = 18; chart_det.height = 8
            
            data_det = Reference(ws_frames, min_col=5, max_col=5, min_row=1, max_row=max_row_frames)
            chart_det.add_data(data_det, titles_from_data=True)
            chart_det.set_categories(cat_frames)
            ws_frames.add_chart(chart_det, "G18") # Рисуем ниже (G18)

            # === ГРАФИКИ ДЛЯ ТЕЛЕМЕТРИИ ===
            ws_telem = writer.sheets['Telemetry']
            max_row_telem = len(df_telem) + 1
            cat_telem = Reference(ws_telem, min_col=1, min_row=2, max_row=max_row_telem) # Ось X (Время)

            def add_telem_chart(title, min_col, max_col, pos):
                ch = LineChart()
                ch.title = title
                ch.x_axis.title = "System Time"
                ch.width = 18; ch.height = 8
                d = Reference(ws_telem, min_col=min_col, max_col=max_col, min_row=1, max_row=max_row_telem)
                ch.add_data(d, titles_from_data=True)
                ch.set_categories(cat_telem)
                ws_telem.add_chart(ch, pos)

            # Колонки: Voltage(6), Current(7), Temp1(8), Temp2(9), Temp3(10), RPM(11)
            add_telem_chart("Voltage (V)", 6, 6, "M2")
            add_telem_chart("Current (A)", 7, 7, "M18")
            add_telem_chart("Temperatures (°C)", 8, 10, "M34")
            add_telem_chart("Real RPM", 11, 11, "M50")

        print(f"[SUCCESS] Report with Interactive Charts saved to: {output_path}")
    except Exception as e:
        print(f"[ERROR] Failed to save Excel file: {e}")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="LiDAR Telemetry & Detection Excel Report Generator")
    
    parser.add_argument("input", help="Path to the .pcap or .pcapng file")
    args = parser.parse_args()
    parser.add_argument("-o", "--output", default=f"{args.input}_report.xlsx", help="Path to the output .xlsx file")
    parser.add_argument("--msop", type=int, default=2368, help="UDP Port for MSOP data (default: 2368)")
    parser.add_argument("--difop", type=int, default=8368, help="UDP Port for DIFOP data (default: 8368)")
    
    parser.add_argument("--roi-min", nargs=3, type=float, default=[-0.61560, -0.27765, -0.00780],
                        metavar=('X', 'Y', 'Z'), help="Min coordinates for Bounding Box")
    parser.add_argument("--roi-max", nargs=3, type=float, default=[0.25920, 0.27765, 0.47000],
                        metavar=('X', 'Y', 'Z'), help="Max coordinates for Bounding Box")
    
    args = parser.parse_args()
    roi_min_arr = np.array(args.roi_min)
    roi_max_arr = np.array(args.roi_max)
    generate_report(args.input, args.output, args.msop, args.difop, roi_min_arr, roi_max_arr)